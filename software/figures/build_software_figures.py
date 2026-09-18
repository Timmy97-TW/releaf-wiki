"""Figures and numbers for the wiki Software page (software/index.html).

    python3 software/figures/build_software_figures.py [--project <path>] [--downloads <path>]

Run from the wiki root. Writes:

  assets/img/software/fig-r1-hydraulics.svg   R1, 260816 pressure to flow rate.xlsx
  assets/img/software/fig-r2-photometer.svg   R2, Inline photometer to biodrop.xlsx
  assets/img/software/fig-r3-run434.svg       R3, out/sentinel_series.json + out/sentinel.json
  assets/img/software/fig-r4-run0902.svg      R4, record_20260906_090304.csv
  assets/img/software/fig-e1..e4-*.svg        MOCK templates for the evidence package (no data)
  software/figures/numbers.json               every number quoted on the page, with its source file

Only numpy and openpyxl. No plotting library, so the colours are the wiki tokens
(assets/css/tokens.css) exactly and the files are plain, diffable SVG.

The four E figures are templates. They carry axes, the acceptance criterion and
a MOCK banner inside the SVG, and they contain no data points, synthetic or
otherwise, so they cannot be cropped into a slide and mistaken for a result.
"""
import argparse, csv, json, math, pathlib, re
import numpy as np
import openpyxl

WIKI = pathlib.Path(__file__).resolve().parents[2]
ap = argparse.ArgumentParser()
ap.add_argument("--project", default=str(pathlib.Path.home() / "Documents/Claude/Projects/2026 iGEM Project Plant Stress"))
ap.add_argument("--downloads", default=str(pathlib.Path.home() / "Downloads"))
args = ap.parse_args()
PROJECT, DL = pathlib.Path(args.project), pathlib.Path(args.downloads)
MODEL_OUT = PROJECT / "Math Model" / "releaf-model" / "out"
OUT = WIKI / "assets" / "img" / "software"
OUT.mkdir(parents=True, exist_ok=True)

# ---- colours: read from tokens.css so the figures cannot drift from the site --
_tok = dict(re.findall(r"--([a-z0-9-]+):\s*(#[0-9a-fA-F]{6})", (WIKI / "assets/css/tokens.css").read_text()))
INK, INK2, INK3, GRID, FAINT = _tok["black"], _tok["gray-600"], _tok["gray-500"], _tok["gray-200"], _tok["gray-50"]
LEAF9, LEAF7, LEAF5, LEAF2, LEAF1 = _tok["leaf-900"], _tok["leaf-700"], _tok["leaf-500"], _tok["leaf-200"], _tok["leaf-100"]
RUST, RUST1 = _tok["rust-700"], _tok["rust-100"]
AMBER, AMBER1 = _tok["amber-700"], _tok["amber-100"]
SLATE, SLATE1 = _tok["slate-700"], _tok["slate-100"]
FONT = "Inter, -apple-system, 'Segoe UI', Helvetica, Arial, sans-serif"

NUM = {}          # everything the page quotes, keyed by result


def rec(key, source, **vals):
    NUM[key] = {"source": source, **{k: (float(f"{float(v):.6g}") if isinstance(v, (float, np.floating)) else v)
                                     for k, v in vals.items()}}


# ------------------------------------------------------------------ drawing --
def esc(t):
    return str(t).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


class Fig:
    def __init__(self, title, desc, w=880, h=440, bg="#ffffff"):
        self.w, self.h, self.o = w, h, []
        fill = f'<rect width="{w}" height="{h}" fill="{bg}"/>' if bg else ""
        self.head = (f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {w} {h}" width="{w}" height="{h}" '
                     f'role="img" font-family="{FONT}"><title>{esc(title)}</title><desc>{esc(desc)}</desc>{fill}')

    def add(self, s):
        self.o.append(s)

    def text(self, x, y, t, size=12, fill=INK2, anchor="start", weight=400, italic=False, rot=None):
        tr = f' transform="rotate({rot} {x:.1f} {y:.1f})"' if rot is not None else ""
        st = ' font-style="italic"' if italic else ""
        self.add(f'<text x="{x:.1f}" y="{y:.1f}" font-size="{size}" fill="{fill}" text-anchor="{anchor}" '
                 f'font-weight="{weight}"{st}{tr}>{esc(t)}</text>')

    def mock(self):
        self.add(f'<rect x="{self.w-178}" y="12" width="166" height="26" rx="3" fill="{AMBER1}" '
                 f'stroke="{AMBER}" stroke-dasharray="4 3"/>')
        self.text(self.w-95, 30, "MOCK · NO DATA YET", 11.5, AMBER, "middle", 700)

    def save(self, name):
        (OUT / name).write_text(self.head + "".join(self.o) + "</svg>\n", encoding="utf-8")
        print("wrote", OUT.relative_to(WIKI) / name)


class Axes:
    def __init__(self, f, box, xlim, ylim, xlabel, ylabel, xticks, yticks, yfmt="{:g}", xfmt="{:g}", ylab_dx=50):
        self.f, (self.L, self.T, self.W, self.H) = f, box
        self.xl, self.yl = xlim, ylim
        L, T, Wd, Hd = box
        for v in yticks:
            y = self.Y(v)
            f.add(f'<line x1="{L}" y1="{y:.1f}" x2="{L+Wd}" y2="{y:.1f}" stroke="{GRID}"/>')
            f.text(L-8, y+4, yfmt.format(v), 11, INK3, "end")
        f.add(f'<line x1="{L}" y1="{T+Hd}" x2="{L+Wd}" y2="{T+Hd}" stroke="{INK3}"/>')
        for v in xticks:
            x = self.X(v)
            f.add(f'<line x1="{x:.1f}" y1="{T+Hd}" x2="{x:.1f}" y2="{T+Hd+4}" stroke="{INK3}"/>')
            f.text(x, T+Hd+18, xfmt.format(v), 11, INK3, "middle")
        if xlabel:
            f.text(L+Wd/2, T+Hd+38, xlabel, 12, INK2, "middle")
        f.text(L-ylab_dx, T+Hd/2, ylabel, 12, INK2, "middle", rot=-90)

    def X(self, v):
        return self.L + (v-self.xl[0])/(self.xl[1]-self.xl[0])*self.W

    def Y(self, v):
        return self.T + self.H - (v-self.yl[0])/(self.yl[1]-self.yl[0])*self.H

    def line(self, xs, ys, color, width=1.8, dash=None):
        pts = " ".join(f"{self.X(x):.1f},{self.Y(y):.1f}" for x, y in zip(xs, ys)
                       if self.xl[0] <= x <= self.xl[1] and np.isfinite(y) and self.yl[0] <= y <= self.yl[1])
        d = f' stroke-dasharray="{dash}"' if dash else ""
        self.f.add(f'<polyline points="{pts}" fill="none" stroke="{color}" stroke-width="{width}" '
                   f'stroke-linejoin="round"{d}/>')

    def dots(self, xs, ys, color, r=3.2, fill=None):
        fill = fill or color
        self.f.add("".join(f'<circle cx="{self.X(x):.1f}" cy="{self.Y(y):.1f}" r="{r}" fill="{fill}" '
                           f'stroke="{color}" stroke-width="1.2"/>' for x, y in zip(xs, ys)
                           if np.isfinite(y) and self.yl[0] <= y <= self.yl[1] and self.xl[0] <= x <= self.xl[1]))

    def band(self, x0, x1, color):
        self.f.add(f'<rect x="{self.X(x0):.1f}" y="{self.T}" width="{self.X(x1)-self.X(x0):.1f}" '
                   f'height="{self.H}" fill="{color}"/>')

    def hband(self, y0, y1, color):
        self.f.add(f'<rect x="{self.L}" y="{self.Y(y1):.1f}" width="{self.W}" height="{self.Y(y0)-self.Y(y1):.1f}" '
                   f'fill="{color}"/>')

    def vline(self, x, color, dash="4 3", width=1.4):
        self.f.add(f'<line x1="{self.X(x):.1f}" y1="{self.T}" x2="{self.X(x):.1f}" y2="{self.T+self.H}" '
                   f'stroke="{color}" stroke-width="{width}" stroke-dasharray="{dash}"/>')

    def hline(self, y, color, dash="4 3", width=1.4):
        self.f.add(f'<line x1="{self.L}" y1="{self.Y(y):.1f}" x2="{self.L+self.W}" y2="{self.Y(y):.1f}" '
                   f'stroke="{color}" stroke-width="{width}" stroke-dasharray="{dash}"/>')


def legend(f, x, y, items):
    for i, (label, color, kind) in enumerate(items):
        yy = y + i*20
        if kind == "dot":
            f.add(f'<circle cx="{x+8}" cy="{yy-4}" r="3.6" fill="{color}"/>')
        elif kind == "ring":
            f.add(f'<circle cx="{x+8}" cy="{yy-4}" r="3.6" fill="#fff" stroke="{color}" stroke-width="1.4"/>')
        elif kind == "dash":
            f.add(f'<line x1="{x}" y1="{yy-4}" x2="{x+16}" y2="{yy-4}" stroke="{color}" stroke-width="2" stroke-dasharray="5 3"/>')
        elif kind == "band":
            f.add(f'<rect x="{x}" y="{yy-10}" width="16" height="11" fill="{color}"/>')
        else:
            f.add(f'<line x1="{x}" y1="{yy-4}" x2="{x+16}" y2="{yy-4}" stroke="{color}" stroke-width="2.2"/>')
        f.text(x+23, yy, label, 11.5, INK2)


# --------------------------------------------------------------- statistics --
def r2(y, yhat):
    y, yhat = np.asarray(y, float), np.asarray(yhat, float)
    return float(1 - np.sum((y-yhat)**2)/np.sum((y-y.mean())**2))


def fit_origin_quadratic(q, dp):
    """dP = aQ + bQ^2 through the origin. Residual SD with 2 fitted parameters."""
    A = np.column_stack([q, q**2])
    coef, *_ = np.linalg.lstsq(A, dp, rcond=None)
    res = dp - A @ coef
    return coef, r2(dp, A @ coef), float(np.sqrt(np.sum(res**2)/(len(dp)-2)))


def linfit(x, y):
    c = np.polyfit(x, y, 1)
    res = y - np.polyval(c, x)
    return c, r2(y, np.polyval(c, x)), float(np.sqrt(np.sum(res**2)/(len(x)-2)))


# ======================================================= R1 · hydraulics =====
def fig_r1():
    src = DL / "260816 pressure to flow rate.xlsx"
    ws = openpyxl.load_workbook(src, data_only=True).worksheets[0]
    rows = [r for r in ws.iter_rows(min_row=3, values_only=True) if isinstance(r[0], (int, float))]
    pump = np.array([r[0] for r in rows], float)
    qj, pj = np.array([r[1] for r in rows], float), np.array([r[3] for r in rows], float)   # jumper only
    qm, pm = np.array([r[4] for r in rows], float), np.array([r[6] for r in rows], float)   # module fitted
    (am, bm), r2m, sdm = fit_origin_quadratic(qm, pm)
    (aj, bj), r2j, sdj = fit_origin_quadratic(qj, pj)
    an, bn = am - aj, bm - bj
    share = (aj*qm + bj*qm**2) / pm                       # rig share of the measured total, at each module point
    q_op = 263.0
    dp_net_op = an*q_op + bn*q_op**2
    dp_tot_op = am*q_op + bm*q_op**2
    # Hagen-Poiseuille for 11 open fibres, 0.5 mm bore radius, 0.60 m, water 22 C (rig_constants.csv)
    mu, L, r, n = 0.955e-3, 0.60, 0.5e-3, 11
    hp = lambda q: 8*mu*L*(q*1e-6/60/n)/(math.pi*r**4)/1e5            # bar
    rec("R1", "~/Downloads/260816 pressure to flow rate.xlsx (sheet 1; columns B,D jumper only; E,G module fitted)",
        n_points=int(len(pump)), pump_pct=[int(p) for p in pump],
        a_module=am, b_module=bm, r2_module=r2m, sd_module_mbar=sdm*1000,
        a_jumper=aj, b_jumper=bj, r2_jumper=r2j, sd_jumper_mbar=sdj*1000,
        a_net=an, b_net=bn, rig_share_min=float(share.min()), rig_share_max=float(share.max()),
        q_operating=q_op, dp_total_operating_bar=dp_tot_op, dp_net_operating_bar=dp_net_op,
        hagen_poiseuille_operating_bar=hp(q_op), net_over_hp_operating=dp_net_op/hp(q_op),
        q_min=float(min(qm.min(), qj.min())), q_max=float(max(qm.max(), qj.max())))

    f = Fig("R1. Pressure drop across the hollow-fibre loop, 16 August 2026",
            "Measured pressure drop against flow with the module fitted and with a jumper only, each fitted with "
            "dP = aQ + bQ^2 through the origin, and the net module curve drawn as the difference of the two fits.")
    ax = Axes(f, (80, 30, 560, 340), (0, 480), (0, 0.5), "Loop flow Q (mL/min)", "Pressure drop ΔP (bar)",
              range(0, 481, 80), [0, .1, .2, .3, .4, .5], "{:.1f}")
    qq = np.linspace(0, 470, 160)
    ax.line(qq, am*qq + bm*qq**2, LEAF7, 2)
    ax.line(qq, aj*qq + bj*qq**2, SLATE, 2)
    ax.line(qq, an*qq + bn*qq**2, RUST, 2, "6 4")
    ax.dots(qm, pm, LEAF7, 3.6)
    ax.dots(qj, pj, SLATE, 3.6, "#fff")
    ax.vline(q_op, INK3, "2 3", 1.2)
    f.text(ax.X(q_op)-6, ax.Y(0.47), f"operating point {q_op:.0f} mL/min", 10.5, INK3, "end")
    x0 = 664
    legend(f, x0, 52, [("Module fitted (total)", LEAF7, "dot"), ("Jumper only (rig)", SLATE, "ring"),
                       ("Net module = difference of fits", RUST, "dash")])
    f.text(x0, 136, "Fit ΔP = aQ + bQ², through origin", 12, INK, weight=650)
    f.text(x0, 158, f"Module fitted  R² {r2m:.4f}", 11.5, INK2)
    f.text(x0, 176, f"   residual SD {sdm*1000:.1f} mbar", 11.5, INK3)
    f.text(x0, 198, f"Jumper only    R² {r2j:.4f}", 11.5, INK2)
    f.text(x0, 216, f"   residual SD {sdj*1000:.1f} mbar", 11.5, INK3)
    f.text(x0, 244, "Net module", 12, RUST, weight=650)
    f.text(x0, 264, f"a = {an*1e4:.3f} × 10⁻⁴ bar·min/mL", 11.5, INK2)
    f.text(x0, 283, f"b = {bn*1e7:.3f} × 10⁻⁷ bar·min²/mL²", 11.5, INK2)
    f.text(x0, 318, "Water, 22 °C, 8 pump settings,", 11, INK3, italic=True)
    f.text(x0, 335, "n = 1 per point. A calibration:", 11, INK3, italic=True)
    f.text(x0, 352, "no point was held out.", 11, INK3, italic=True)
    f.save("fig-r1-hydraulics.svg")


# ======================================================= R2 · photometer =====
def fig_r2():
    src = DL / "Inline photometer to biodrop.xlsx"
    rows = [r for r in openpyxl.load_workbook(src, data_only=True).worksheets[0].iter_rows(min_row=2, values_only=True)
            if isinstance(r[0], (int, float))]
    m = np.array([r[0] for r in rows], float)
    ours = np.array([r[1] if isinstance(r[1], (int, float)) else np.nan for r in rows], float)
    bd = np.array([r[2] if isinstance(r[2], (int, float)) else np.nan for r in rows], float)
    # agreement with the BioDrop where the BioDrop reads below 1.2
    k = ~np.isnan(bd) & (bd < 1.2)
    cab, r2ab, sdab = linfit(bd[k], ours[k])
    diff = ours[k] - bd[k]
    bias, loa = float(diff.mean()), 1.96*float(diff.std(ddof=1))
    # linearity against dose, above 30 mg
    # scored over the range both instruments cover (30 < m <= 46 mg), and ours alone to 60 mg
    kb = ~np.isnan(bd) & (m > 30)
    ko = ~np.isnan(ours) & (m > 30) & (m <= m[kb].max())
    co, r2o, sdo = linfit(m[ko], ours[ko])
    k60 = ~np.isnan(ours) & (m > 30)
    c60, r2o60, sdo60 = linfit(m[k60], ours[k60])
    cb, r2b, sdb = linfit(m[kb], bd[kb])
    kl = ~np.isnan(ours) & (m <= 30)
    cl, r2l, sdl = linfit(m[kl], ours[kl])
    rec("R2", "~/Downloads/Inline photometer to biodrop.xlsx (sheet 1; TiO2 mg, in-line photometer, BioDrop)",
        n_ours=int((~np.isnan(ours)).sum()), n_biodrop=int((~np.isnan(bd)).sum()),
        mass_min=float(m.min()), mass_max=float(m.max()), biodrop_max_mass=float(m[~np.isnan(bd)].max()),
        agree_n=int(k.sum()), agree_slope=cab[0], agree_intercept=cab[1], agree_r2=r2ab, agree_sd=sdab,
        agree_bias=bias, agree_loa_low=bias-loa, agree_loa_high=bias+loa,
        ours_le30_r2=r2l, ours_le30_slope=cl[0],
        ours_gt30_n=int(ko.sum()), ours_gt30_r2=r2o, ours_gt30_sd=sdo,
        ours_30_60_n=int(k60.sum()), ours_30_60_r2=r2o60, ours_30_60_sd=sdo60,
        biodrop_gt30_n=int(kb.sum()), biodrop_gt30_r2=r2b, biodrop_gt30_sd=sdb)

    f = Fig("R2. In-line photometer and BioDrop against TiO2 dose, 23 August 2026",
            "Optical density reported by the in-line photometer and by a BioDrop for cumulative TiO2 additions "
            "into 400 mL, with the region above 30 mg shaded and a straight line fitted to each instrument there.")
    ax = Axes(f, (80, 30, 560, 340), (0, 62), (0, 3.0), "Cumulative TiO₂ added (mg in 400 mL)", "Reported OD",
              range(0, 61, 10), [0, .5, 1, 1.5, 2, 2.5, 3], "{:g}")
    ax.band(30, 62, FAINT)
    f.add(f'<line x1="{ax.X(30):.1f}" y1="{ax.T}" x2="{ax.X(30):.1f}" y2="{ax.T+ax.H}" stroke="{GRID}" stroke-width="1.5"/>')
    f.text(ax.X(31), ax.T+16, "above 30 mg: linearity scored against dose", 10.5, INK3, italic=True)
    xx = np.linspace(30, 46, 20)
    ax.line(xx, np.polyval(co, xx), LEAF9, 1.4, "5 3")
    xb = np.linspace(30, 46, 20)
    ax.line(xb, np.polyval(cb, xb), SLATE, 1.4, "5 3")
    ax.dots(m, bd, SLATE, 3.2, "#fff")
    ax.dots(m, ours, LEAF7, 3.2)
    x0 = 664
    legend(f, x0, 52, [("In-line photometer (ours)", LEAF7, "dot"), ("BioDrop", SLATE, "ring"),
                       ("Line fitted, 31–46 mg", INK3, "dash")])
    f.text(x0, 136, "31–46 mg, fit against dose", 12, INK, weight=650)
    f.text(x0, 158, f"Ours     R² {r2o:.4f}, SD {sdo:.3f}", 11.5, LEAF9)
    f.text(x0, 176, f"BioDrop  R² {r2b:.4f}, SD {sdb:.3f}", 11.5, SLATE)
    f.text(x0, 204, "Below BioDrop OD 1.2 (n = %d)" % k.sum(), 12, INK, weight=650)
    f.text(x0, 226, f"ours = {cab[0]:.3f} × BioDrop − {abs(cab[1]):.3f}", 11.5, INK2)
    f.text(x0, 244, f"R² {r2ab:.3f}, residual SD {sdab:.3f}", 11.5, INK2)
    f.text(x0, 284, "One reading per point, no replicates;", 11, INK3, italic=True)
    f.text(x0, 301, "additions cumulative. TiO₂ tests", 11, INK3, italic=True)
    f.text(x0, 318, "linearity, not cell number.", 11, INK3, italic=True)
    f.save("fig-r2-photometer.svg")


# ============================================================ R3 · 434 h =====
def fig_r3():
    s = json.load(open(MODEL_OUT / "sentinel_series.json"))
    meta = json.load(open(MODEL_OUT / "sentinel.json"))
    lt = meta["lead_time"]
    t, od, ref = (np.array(s[k], float) for k in ("t_h", "od", "ref_lux"))
    alarm, unusable = lt["sensor_alarm_autoencoder_h"]["value"], lt["od_unusable_from_h"]["value"]
    rec("R3", "Math Model/releaf-model/out/sentinel_series.json and out/sentinel.json",
        n_series=int(len(t)), n_rows_valid=meta["run"]["n_rows_valid"]["value"], span_h=meta["run"]["span_h"]["value"],
        interval_min=meta["run"]["interval_min"]["value"], raw_log=meta["run"]["source"]["value"],
        n_train=meta["split"]["n_train"]["value"],
        alarm_ae_h=alarm, alarm_pca_h=lt["sensor_alarm_pca_h"]["value"], od_unusable_h=unusable,
        lead_ae_h=lt["lead_autoencoder_h"]["value"], lead_pca_h=lt["lead_pca_h"]["value"],
        far_unseen_quiet_ae=meta["benchmark"]["sensor"]["methods"]["ae"]["false_alarm_rate_unseen_quiet_160_220h"],
        process_window_fired_ae=meta["benchmark"]["sensor"]["methods"]["ae"]["process_120_160_fired_frac"],
        ref_final_lux=lt["reference_channel_final"]["value"])

    f = Fig("R3. The 434-hour run: OD600 and the photometer reference channel",
            "Two panels over 434 hours at 22 C. Top: in-line OD600. Bottom: the photometer reference channel. "
            "The process event (120-160 h) and the instrument event (215-300 h) are shaded; the trust-layer alarm "
            "at 216.3 h and the start of unusable OD at 243.0 h are marked.", 880, 480)
    a1 = Axes(f, (80, 36, 580, 220), (0, 440), (-0.2, 2.0), "", "OD600 (in-line)", [], [0, .5, 1, 1.5, 2], "{:g}", ylab_dx=46)
    a2 = Axes(f, (80, 290, 580, 120), (0, 440), (0, 30), "Elapsed time (h)", "Reference (lux)",
              range(0, 441, 40), [0, 10, 20, 30], "{:g}", ylab_dx=46)
    for a in (a1, a2):
        a.band(120, 160, SLATE1)
        a.band(215, 300, RUST1)
    a1.dots(t, od, LEAF7, 1.1)
    a2.line(t, ref, SLATE, 1.2)
    for a in (a1, a2):
        a.vline(alarm, RUST, "5 3", 1.6)
        a.vline(unusable, INK, "2 3", 1.4)
    f.text(a1.X(140), a1.T+14, "process", 10.5, SLATE, "middle", 650)
    f.text(a1.X(140), a1.T+28, "event", 10.5, SLATE, "middle", 650)
    f.text(a1.X(300)-6, a1.T+14, "instrument event", 10.5, RUST, "end", 650)
    x0 = 684
    f.add(f'<line x1="{x0}" y1="60" x2="{x0+18}" y2="60" stroke="{RUST}" stroke-width="1.6" stroke-dasharray="5 3"/>')
    f.text(x0+26, 64, f"{alarm:.1f} h  sensor alarm", 11.5, RUST, weight=650)
    f.add(f'<line x1="{x0}" y1="84" x2="{x0+18}" y2="84" stroke="{INK}" stroke-width="1.4" stroke-dasharray="2 3"/>')
    f.text(x0+26, 88, f"{unusable:.1f} h  OD unusable", 11.5, INK, weight=650)
    f.text(x0, 124, f"Lead time {lt['lead_autoencoder_h']['value']:.1f} h", 16, LEAF9, weight=700)
    f.text(x0, 154, "Trained on t < 120 h only;", 11.5, INK2)
    f.text(x0, 172, "limit = 99.5th percentile of", 11.5, INK2)
    f.text(x0, 190, "the training score, fixed", 11.5, INK2)
    f.text(x0, 208, "before scoring.", 11.5, INK2)
    f.text(x0, 300, "Reference falls while OD", 11, INK3, italic=True)
    f.text(x0, 317, "scatter grows: the instrument", 11, INK3, italic=True)
    f.text(x0, 334, "event. In 120–160 h the", 11, INK3, italic=True)
    f.text(x0, 351, "reference is steady.", 11, INK3, italic=True)
    f.text(x0, 380, "Two events, one run:", 11, INK3, italic=True)
    f.text(x0, 397, "no error rate is reportable.", 11, INK3, italic=True)
    f.save("fig-r3-run434.svg")


# ========================================================== R4 · 102 h =====
def fig_r4():
    src = DL / "record_20260906_090304.csv"
    t, od, ref, tmp = [], [], [], []
    n_all = 0
    with open(src, newline="") as fh:
        for r in csv.DictReader(fh):
            n_all += 1
            try:
                tt = float(r["elapsed_s"])/3600
                if r["pressure_tmp_bar"]:
                    tmp.append((tt, float(r["pressure_tmp_bar"])))
                if r["od600_valid"].strip().lower() != "true":
                    continue
                if float(r["od600_bubble_pct"] or 0) > 20:
                    continue
                t.append(tt); od.append(float(r["od600_au"])); ref.append(float(r["od600_lux_ref"]))
            except (ValueError, KeyError):
                continue
    t, od, ref = np.array(t), np.array(od), np.array(ref)
    th, oh = [], []
    for a in np.arange(0, math.ceil(t.max()), 1.0):
        k = (t >= a) & (t < a+1)
        if k.sum():
            th.append(a+0.5); oh.append(float(np.median(od[k])))
    th, oh = np.array(th), np.array(oh)
    cl, r2l, _ = linfit(th, oh)
    ce = np.polyfit(th, np.log(oh), 1)
    r2e = r2(oh, np.exp(np.polyval(ce, th)))
    first = th <= 20; last = th >= th.max()-20
    s_first = np.polyfit(th[first], oh[first], 1)[0]; s_last = np.polyfit(th[last], oh[last], 1)[0]
    # reference-channel step at the 28.7 / 30.8 h interventions
    r_pre = float(np.median(ref[(t >= 14) & (t < 26)])); r_post = float(np.median(ref[(t >= 33) & (t < 45)]))
    step = r_post/r_pre - 1
    od_offset = 5*math.log10(r_post/r_pre)            # OD = 5 log10(k ref/sample)
    tt = np.array(tmp)
    ctmp = np.polyfit(tt[:, 0], tt[:, 1], 1)
    quiet = tt[(tt[:, 0] >= 60) & (tt[:, 0] <= 100), 1]
    rec("R4", "~/Downloads/record_20260906_090304.csv (30 s medians; od600_valid true, od600_bubble_pct <= 20)",
        n_rows=n_all, n_kept=int(len(t)), n_hours=int(len(th)), span_h=float(t.max()),
        od_first_hour=float(oh[0]), od_last_hour=float(oh[-1]),
        lin_slope=cl[0], lin_intercept=cl[1], lin_r2=r2l, exp_mu=ce[0], exp_r2=r2e,
        slope_first20=s_first, slope_last20=s_last,
        ref_pre_lux=r_pre, ref_post_lux=r_post, ref_step_pct=step*100, od_offset_from_step=od_offset,
        tmp_range_bar=float(tt[:, 1].max()-tt[:, 1].min()), tmp_median_bar=float(np.median(tt[:, 1])),
        tmp_slope_mbar_per_h=ctmp[0]*1000, tmp_drift_bound_mbar=ctmp[0]*1000*float(t.max()),
        tmp_quiet_sd_mbar=float(quiet.std(ddof=1))*1000)

    f = Fig("R4. The 102-hour perfusion run, 2 to 6 September 2026",
            "Hourly medians of in-line OD600 over 102 hours after removing invalid and bubble-affected rows, "
            "with a linear and an exponential fit.")
    ax = Axes(f, (80, 30, 560, 340), (0, 105), (0, 4.2), "Elapsed time (h)", "OD600 (in-line, hourly median)",
              range(0, 101, 20), [0, 1, 2, 3, 4], "{:g}")
    ax.dots(t[::30], od[::30], LEAF2, 1.2)
    xx = np.linspace(0, 104, 120)
    ax.line(xx, np.exp(np.polyval(ce, xx)), RUST, 1.8, "6 4")
    ax.line(xx, np.polyval(cl, xx), LEAF9, 2)
    ax.dots(th, oh, LEAF7, 2.8)
    for x in (28.7, 30.8):
        ax.vline(x, INK3, "2 3", 1)
    f.text(ax.X(31.5), ax.Y(4.0), "operator interventions, 28.7 and 30.8 h", 10.5, INK3, italic=True)
    x0 = 664
    legend(f, x0, 52, [("30 s windows (every 30th)", LEAF2, "dot"), ("Hourly median", LEAF7, "dot"),
                       ("Linear fit", LEAF9, "line"), ("Exponential fit", RUST, "dash")])
    f.text(x0, 150, f"Linear   R² {r2l:.3f}", 12, LEAF9, weight=650)
    f.text(x0, 169, f"   slope {cl[0]:.4f} OD/h", 11.5, INK2)
    f.text(x0, 193, f"Exponential   R² {r2e:.3f}", 12, RUST, weight=650)
    f.text(x0, 212, f"   µ {ce[0]:.4f} /h", 11.5, INK2)
    f.text(x0, 244, f"First 20 h {s_first:.3f} OD/h;", 11.5, INK2)
    f.text(x0, 262, f"last 20 h {s_last:.3f} OD/h.", 11.5, INK2)
    f.text(x0, 300, "In-line OD was not checked", 11, INK3, italic=True)
    f.text(x0, 317, f"off-line. A −{abs(step)*100:.2f} % reference", 11, INK3, italic=True)
    f.text(x0, 334, f"step at 29 h shifts later OD", 11, INK3, italic=True)
    f.text(x0, 351, f"by −{abs(od_offset):.3f}.", 11, INK3, italic=True)
    f.save("fig-r4-run0902.svg")


# ============================================ E1–E4 · MOCK templates ========
# 16:9, transparent, so the .fig--pending hatching shows through. No data.
MW, MH = 880, 495


def mock_note(f, x0, y0, head, lines, crit):
    f.text(x0, y0, head, 12, INK, weight=650)
    for i, s in enumerate(lines):
        f.text(x0, y0+24+i*19, s, 11.5, INK2)
    yc = y0+24+len(lines)*19+18
    f.text(x0, yc, "Acceptance (fixed in advance)", 11.5, LEAF9, weight=700)
    for i, s in enumerate(crit):
        f.text(x0, yc+19+i*18, s, 11.5, LEAF9)


def fig_e1():
    f = Fig("MOCK. E1: link loss and latency (DT-1)", "Template, no data. Left: a commanded pump step and the "
            "PT-01 response against time from command. Right: what must be reported.", MW, MH, bg=None)
    f.mock()
    ax = Axes(f, (80, 56, 520, 360), (-2, 10), (0, 8), "Time from SET sent (s)", "PT-01 (psi)",
              range(-2, 11, 2), [0, 2, 4, 6, 8], "{:g}")
    ax.vline(0, INK, "4 3")
    f.text(ax.X(0)+6, ax.T+16, "SET P01 sent", 11, INK)
    f.text(ax.X(4), ax.Y(4), "20 steps per direction, overlaid", 12, INK3, "middle", italic=True)
    mock_note(f, 630, 90, "Report", ["PV lines expected vs received", "loss = missing counters N / expected",
                                     "longest gap between PV lines (s)", "SET → ACK: median, IQR (ms)",
                                     "SET → 63 % of ΔPT-01: median, IQR (s)"],
              ["loss < 1 %; no gap > 5 s", "fail-safe stop ≤ 6 s in 3 of 3", "median SET → 63 % ≤ 10 s"])
    f.save("fig-e1-link.svg")


def fig_e2():
    f = Fig("MOCK. E2: held-out parity (DT-3 and DT-5)", "Template, no data. Predicted against measured on a "
            "session the model was not fitted to, with the 1:1 line.", MW, MH, bg=None)
    f.mock()
    ax = Axes(f, (80, 56, 360, 360), (0, 0.5), (0, 0.5), "Measured lumen ΔP (bar), held-out session",
              "Predicted lumen ΔP (bar), no refit", [0, .1, .2, .3, .4, .5], [0, .1, .2, .3, .4, .5], "{:.1f}", "{:.1f}")
    ax.line([0, .5], [0, .5], INK3, 1.2, "4 3")
    f.text(ax.X(.42)+6, ax.Y(.42)+4, "1 : 1", 11, INK3)
    mock_note(f, 480, 90, "Report, held-out points only", ["R²p", "RMSEP (bar)", "bias = mean(pred − meas)",
                                                            "RPD = SD(measured) / RMSEP",
                                                            "Same form for DT-5 with OD600 on both axes"],
              ["DT-3: RPD ≥ 3 and RMSEP ≤ 20 mbar", "DT-5: RPD ≥ 3 at a 4 h horizon"])
    f.save("fig-e2-parity.svg")


def fig_e3():
    f = Fig("MOCK. E3: fouling forecast error against horizon (DT-4)", "Template, no data. Forecast error of "
            "net module pressure drop as a function of horizon, against the measurement noise floor.", MW, MH, bg=None)
    f.mock()
    ax = Axes(f, (80, 56, 520, 360), (0, 13), (0, 30), "Forecast horizon (h)", "RMSE of ΔP_net forecast (mbar)",
              [1, 3, 6, 12], [0, 10, 20, 30], "{:g}")
    ax.hband(0, 10, FAINT)
    ax.hline(10, INK3, "4 3", 1.2)
    f.text(ax.X(12.8), ax.Y(10)-6, "10 mbar: residual SD of R1 (module fitted)", 10.5, INK3, "end", italic=True)
    mock_note(f, 630, 90, "Report", ["rolling origins every 1 h, 24–36 h", "fit on 0 to origin; no later data",
                                     "RMSE at 1, 3, 6 and 12 h ahead", "same for persistence forecast"],
              ["12 h RMSE ≤ 10 mbar and below", "persistence, on an informative run"])
    f.save("fig-e3-horizon.svg")


def fig_e4():
    f = Fig("MOCK. E4: fixed schedule against model-triggered back-flush (DT-7)", "Template, no data. Net module "
            "resistance over time under a fixed schedule and under model-triggered back-flushing.", MW, MH, bg=None)
    f.mock()
    ax = Axes(f, (80, 56, 520, 360), (0, 24), (0.9, 2.3), "Elapsed time (h)", "R / R₀ (module resistance)",
              range(0, 25, 4), [1, 1.25, 1.5, 1.75, 2, 2.25], "{:.2f}")
    ax.hline(2.0, RUST, "5 3", 1.4)
    f.text(ax.X(1), ax.Y(2.0)-6, "clean-before limit, 2 × baseline (rig_constants.csv)", 10.5, RUST)
    mock_note(f, 630, 90, "Report, per arm (4 runs each)", ["time above 1.5 × R₀ (h)", "number of Clean cycles",
                                                            "mean R / R₀ over 0–24 h", "difference with 95 % CI"],
              ["closed loop lower on time", "above 1.5 × R₀, paired 95 % CI", "excludes 0, no more Cleans"])
    f.save("fig-e4-closedloop.svg")


if __name__ == "__main__":
    fig_r1(); fig_r2(); fig_r3(); fig_r4()
    fig_e1(); fig_e2(); fig_e3(); fig_e4()
    out = WIKI / "software" / "figures" / "numbers.json"
    out.write_text(json.dumps(NUM, indent=1, ensure_ascii=False) + "\n", encoding="utf-8")
    print("wrote", out.relative_to(WIKI))
